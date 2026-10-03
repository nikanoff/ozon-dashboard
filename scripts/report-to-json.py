#!/usr/bin/env python
"""Extracts the table from an Ozon realization report into JSON.

The report is a formal document: several title rows, a two-level header, then one row per
realized item, then a block of summary lines. This pulls out the table and the summaries so
the numbers can be compared with the Seller API.

Usage:
    python scripts/report-to-json.py <report.xlsx> <out.json>
"""

import json
import re
import sys

from openpyxl import load_workbook

# Column positions in the item table, as the document lays them out.
COLUMNS = {
    "n": 0,
    "name": 1,
    "offer_id": 2,
    "sku": 3,
    "barcode": 4,
    "realized": 5,
    "realized_loyalty": 6,
    "realized_qty": 7,
    "realized_price": 8,
    "returned": 9,
    "returned_loyalty": 10,
    "returned_qty": 11,
    "returned_price": 12,
}


def text(value):
    return "" if value is None else str(value).strip()


def number(value):
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    cleaned = text(value).replace("\u00a0", "").replace(" ", "").replace(",", ".")
    match = re.search(r"-?\d+(?:\.\d+)?", cleaned)
    return float(match.group()) if match else None


def main():
    source, target = sys.argv[1], sys.argv[2]

    workbook = load_workbook(source, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    rows = [[cell for cell in row] for row in sheet.iter_rows(values_only=True)]
    workbook.close()

    cells = [[text(value) for value in row] for row in rows]

    # The item table starts under a header that names the article column; the summaries sit
    # below a row that starts with "Итого реализовано".
    header_at = None
    summary_at = None
    for index, row in enumerate(cells):
        joined = " ".join(row)
        if header_at is None and "Артикул" in joined and "SKU" in joined:
            header_at = index
        if summary_at is None and "Итого реализовано" in joined:
            summary_at = index

    if header_at is None:
        raise SystemExit("Не нашёл шапку таблицы (столбец «Артикул»).")

    end = summary_at if summary_at is not None else len(rows)
    items = []
    for row in rows[header_at + 2 : end]:
        first = number(row[COLUMNS["n"]])
        sku = text(row[COLUMNS["sku"]])
        if first is None or not sku:
            continue
        items.append(
            {
                "n": int(first),
                "name": text(row[COLUMNS["name"]]),
                "offer_id": text(row[COLUMNS["offer_id"]]),
                "sku": sku,
                "barcode": text(row[COLUMNS["barcode"]]),
                "realized": number(row[COLUMNS["realized"]]) or 0.0,
                "realized_loyalty": number(row[COLUMNS["realized_loyalty"]]) or 0.0,
                "realized_qty": int(number(row[COLUMNS["realized_qty"]]) or 0),
                "realized_price": number(row[COLUMNS["realized_price"]]),
                "returned": number(row[COLUMNS["returned"]]) or 0.0,
                "returned_loyalty": number(row[COLUMNS["returned_loyalty"]]) or 0.0,
                "returned_qty": int(number(row[COLUMNS["returned_qty"]]) or 0),
                "returned_price": number(row[COLUMNS["returned_price"]]),
            }
        )

    # Summary lines look like `Итого реализовано (за вычетом возвратов) (руб.): | 123456.78`.
    totals = {}
    for row in cells:
        for index, cell in enumerate(row):
            if not cell:
                continue
            if "Итого реализовано" in cell:
                totals["net_realized"] = number(row[index + 1]) if index + 1 < len(row) else None
            elif "Всего выплат от партнёров" in cell:
                totals["net_loyalty"] = number(row[index + 1]) if index + 1 < len(row) else None

    title = next((cell for row in cells[:6] for cell in row if "Отчет о реализации" in cell), "")
    period = next((cell for row in cells[:6] for cell in row if "Реализация товаров за период" in cell), "")
    seller = next(
        (
            cell
            for row in cells
            for cell in row
            if cell.startswith("ИП ") or cell.startswith("ООО ") and "Интернет Решения" not in cell
        ),
        "",
    )

    number_match = re.search(r"№\s*(\d+)", title)
    period_match = re.search(r"с\s*([\d.]+)\s*по\s*([\d.]+)", period)

    result = {
        "source": source,
        "report_number": number_match.group(1) if number_match else None,
        "period": {
            "from": period_match.group(1) if period_match else None,
            "to": period_match.group(2) if period_match else None,
        },
        "seller": seller,
        "items": items,
        "totals": totals,
        "sums": {
            "realized": round(sum(item["realized"] for item in items), 2),
            "realized_loyalty": round(sum(item["realized_loyalty"] for item in items), 2),
            "realized_qty": sum(item["realized_qty"] for item in items),
            "returned": round(sum(item["returned"] for item in items), 2),
            "returned_loyalty": round(sum(item["returned_loyalty"] for item in items), 2),
            "returned_qty": sum(item["returned_qty"] for item in items),
        },
    }

    with open(target, "w", encoding="utf-8") as handle:
        json.dump(result, handle, ensure_ascii=False, indent=1)

    print(
        f"строк={len(items)} реализовано={result['sums']['realized']} "
        f"возвращено={result['sums']['returned']} "
        f"нетто={totals.get('net_realized')} единиц={result['sums']['realized_qty']}"
    )


if __name__ == "__main__":
    main()
